"""Generate the submission Word document."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def p(text):
    doc.add_paragraph(text)


def bullet(text):
    doc.add_paragraph(text, style='List Bullet')


def code(label, text):
    doc.add_paragraph(label).runs[0].italic = True
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)


# ── Title ─────────────────────────────────────────────────────────────────────
title = doc.add_paragraph()
run = title.add_run('Test Automation Framework')
run.bold = True
run.font.size = Pt(18)

doc.add_paragraph('Senior QE Take-Home Assessment — Submission')
doc.add_paragraph()

# ── What was built ────────────────────────────────────────────────────────────
h1('What Was Built and Why')
p(
    'The assessment asks for a test automation framework covering a fintech microservices '
    'system with a User Service, Transaction Service, and Notification Service. No live '
    'application was provided to test against, so the first step was building a simple '
    'mock backend — a small Node.js server that mimics the User and Transaction APIs '
    'described in the brief. This gave the tests something real to run against instead of '
    'pretending a service exists.'
)
p(
    'The test framework itself uses Playwright, a modern open-source testing tool maintained '
    'by Microsoft. It was chosen because it handles both API testing (sending HTTP requests '
    'and checking the responses) and browser-based UI testing in the same project. That '
    'meant one tool, one config file, and one report instead of managing two separate setups.'
)
p('The project is written in TypeScript, which adds type checking on top of JavaScript '
  'and catches a class of errors before the tests even run.')

# ── How to run ────────────────────────────────────────────────────────────────
h1('How to Run the Tests')
p('From a terminal in the project folder, run these three commands in order:')
code('Step 1 — Install dependencies (done once):', 'npm install\nnpx playwright install chromium')
code('Step 2 — Run the full test suite:', 'npm test')
p(
    'That is everything. Playwright automatically starts the mock API server before the '
    'tests begin and shuts it down when they finish. There is no need to open a separate '
    'terminal or start anything manually.'
)
code('To run only the API tests or only the UI tests:', 'npm run test:api\nnpm run test:ui')
code('To open the HTML report in a browser after a run:', 'npm run report')

# ── 1. API Test Suite ─────────────────────────────────────────────────────────
h1('1. API Test Suite')
p(
    'The API tests check the HTTP endpoints directly — they send a request and verify the '
    'response. No browser is involved. This is the fastest and most reliable layer of the '
    'testing pyramid because there is no UI rendering in the way.'
)
p('The four endpoints covered are:')
bullet('POST /api/users — create a new user account')
bullet('GET /api/users/:id — look up a user by their ID')
bullet('POST /api/transactions — submit a new transaction (transfer, deposit, withdrawal)')
bullet('GET /api/transactions/:userId — retrieve all transactions belonging to a user')

h2('CRUD Operations')
p(
    'These tests verify the basic happy path: when you send a valid request, you get back '
    'the expected data. For example, after creating a user the response should contain the '
    'name, email, account type, and a generated ID.'
)

h2('Error Scenarios')
p(
    'A well-built API should reject bad requests with a clear error. These tests confirm that '
    'behaviour:'
)
bullet('Trying to register with an email that already exists returns a 409 (Conflict)')
bullet('Creating a transaction for a user ID that does not exist returns a 404 (Not Found)')
bullet('Sending a transfer without providing a recipient ID returns a 400 (Bad Request)')

h2('Validation Tests')
p(
    'These check that the API enforces its own rules on the data it accepts. If a field is '
    'wrong or missing, the API should say so clearly instead of failing silently.'
)
bullet('Missing name, badly formatted email, or unsupported account type — all return 400 with a description of the problem')
bullet('A transaction amount of zero or a negative number — rejected with 400')

h2('Authentication Tests')
p(
    'The API uses a simple key-based authentication system (an x-api-key header). These '
    'tests confirm the API does not let requests through without it.'
)
bullet('No key at all — returns 401 (Unauthorized)')
bullet('A made-up or incorrect key — returns 403 (Forbidden)')

h2('Code Examples')
p('The tests are structured to be readable and easy to maintain. A few examples:')

code(
    'Example 1 — Happy path: create a user',
    "test('creates a user successfully', async () => {\n"
    "  const payload = buildUserPayload();\n"
    "  const response = await api.createUser(payload);\n"
    "  const user = await expectSuccess(response, 201);\n"
    "  expect(user.name).toBe(payload.name);\n"
    "  expect(user.email).toBe(payload.email);\n"
    "});"
)
p('What this sends to the API:')
code(
    'Request:',
    'POST /api/users\n'
    'x-api-key: dev-api-key-123\n\n'
    '{ "name": "Jane Smith", "email": "jane_001@example.com", "accountType": "premium" }'
)
p('What the API returns:')
code(
    'Response (201 Created):',
    '{ "id": "1", "name": "Jane Smith", "email": "jane_001@example.com",\n'
    '  "accountType": "premium", "createdAt": "2026-01-01T09:00:00.000Z" }'
)

doc.add_paragraph()

code(
    'Example 2 — Validation: reject an invalid account type',
    "test('rejects invalid accountType with 400', async () => {\n"
    "  const payload = buildUserPayload({ accountType: 'vip' });\n"
    "  const response = await api.createUser(payload);\n"
    "  await expectError(response, 400, 'accountType');\n"
    "});"
)
p('The API responds with:')
code(
    'Response (400 Bad Request):',
    '{ "error": "Validation failed",\n'
    '  "details": ["accountType must be one of: basic, premium"] }'
)

doc.add_paragraph()

code(
    'Example 3 — Error handling: transaction for a user that does not exist',
    "test('returns 404 for unknown userId', async () => {\n"
    "  const payload = buildTransactionPayload({ userId: '9999' });\n"
    "  const response = await api.createTransaction(payload);\n"
    "  await expectError(response, 404);\n"
    "});"
)
p('The API responds with:')
code(
    'Response (404 Not Found):',
    '{ "error": "User 9999 not found" }'
)

# ── 2. UI Test Suite ──────────────────────────────────────────────────────────
h1('2. UI Test Suite')
p(
    'The UI tests open a real browser, load the pages, click buttons, and check what '
    'appears on screen — just like a user would. Since no frontend application was '
    'provided, two simple HTML pages were built: one for user registration and one for '
    'creating a transaction. These pages call the same mock API the API tests use.'
)

h2('User Registration Flow')
bullet('Fill in name, email, and account type, click Submit — a confirmation message should appear with the new account ID')
bullet('Try to register with an email that is already in the system — the error from the API should show on the page')
bullet('Try submitting without a name — the browser should stop the form and highlight the missing field before the request is even made')

h2('Transaction Creation Flow')
bullet('Fill in a valid user ID, amount, and recipient, submit a transfer — confirmation message shows the transaction ID and status')
bullet('Use a user ID that does not exist — the error message from the API should appear on screen')

h2('Screenshots on Failure')
p(
    'Playwright is configured to take a screenshot automatically any time a UI test fails. '
    'This means the HTML report after a run shows exactly what was on the screen when '
    'something went wrong, which speeds up debugging significantly.'
)

# ── 3. Utilities ──────────────────────────────────────────────────────────────
h1('3. Test Utilities')
p(
    'Good test utilities reduce repetition and make the tests easier to read. '
    'Here is what was built:'
)

h2('Test Data Factories (factories.ts)')
p(
    'Instead of writing out a full valid user or transaction object in every single test, '
    'the factories create one with sensible defaults. A test only has to specify the '
    'fields it cares about:'
)
code(
    'Example — override just the accountType:',
    "const payload = buildUserPayload({ accountType: 'basic' });\n"
    "// The name, email, and other fields are filled in automatically"
)
p(
    'Emails are generated with a unique counter and timestamp per run, so two tests '
    'never accidentally collide on the same email address.'
)

h2('API Client Wrapper (apiClient.ts)')
p(
    'This wraps the Playwright request object so every test does not have to repeat '
    'the base URL and auth header. Tests call api.createUser() rather than '
    'constructing the full fetch call each time.'
)

h2('Custom Assertions (assertions.ts)')
p(
    'expectSuccess() and expectError() are short helpers that check the HTTP status code '
    'and return the response body. They make the test code read more like plain English '
    'and keep the error messages informative when something fails.'
)

h2('Environment Config (config/env.ts)')
p(
    'The base URL and API key come from config/environments/dev.json. '
    'Both can be overridden by environment variables, so the same test suite '
    'could be pointed at a staging or production service without changing any code.'
)

# ── 4. Reporting ──────────────────────────────────────────────────────────────
h1('4. Reporting')
p('Three output formats are generated every time the tests run:')
bullet(
    'HTML report (playwright-report/index.html) — opens in any browser. '
    'Shows each test with pass/fail, how long it took, the full error message and stack '
    'trace on failure, and a video replay for any browser test that failed.'
)
bullet(
    'JSON results (test-results/results.json) — a structured summary that can be '
    'consumed by a CI system or a dashboard.'
)
bullet(
    'API response logs (test-results/api-logs/) — each logged API call writes a '
    'separate JSON file containing the status code, headers, and body. Useful when '
    'a test fails and you want to see exactly what the API returned.'
)

# ── Folder structure ──────────────────────────────────────────────────────────
h1('Project Folder Structure')
p(
    'Below is a walkthrough of every folder and file created in the project and '
    'what each one is for.'
)

h2('config/')
p(
    'Holds the environment configuration. There is one sub-folder called environments/ '
    'which contains a dev.json file. That file stores the API base URL and the API key '
    'used when running the tests locally. The env.ts file reads that JSON and makes the '
    'values available to the rest of the project. If someone wanted to point the tests at '
    'a different server they would only need to change dev.json, not hunt through the test files.'
)
bullet('config/env.ts — reads and exports the config values')
bullet('config/environments/dev.json — base URL and API key for the local dev environment')

h2('mock-server/')
p(
    'This folder contains the fake API that the tests run against. Since the assessment '
    'did not provide a live backend, a small Node.js server was built here to act as one. '
    'It has all four endpoints from the brief and keeps its data in memory (no database), '
    'which keeps it simple and fast for testing purposes.'
)
bullet('mock-server/server.ts — the Express server: routes, request validation, auth check, error responses')
bullet('mock-server/store.ts — the in-memory data store: holds users and transactions while the server is running, resets between tests')

h2('mock-frontend/')
p(
    'The assessment asks for UI tests but provides no frontend to test against. '
    'This folder has two plain HTML pages — one for registering a user and one for '
    'creating a transaction. They are intentionally simple: no React, no framework, '
    'just HTML forms that call the mock API. Playwright opens these pages in a real '
    'browser and drives them the same way a user would.'
)
bullet('mock-frontend/register.html — user registration form')
bullet('mock-frontend/transaction.html — transaction creation form')
bullet('mock-frontend/register.js — handles the form submit, calls the API, shows success/error message')
bullet('mock-frontend/transaction.js — same for the transaction form')
bullet('mock-frontend/config.js — sets the API base URL for the frontend pages')
bullet('mock-frontend/styles.css — basic styling so the pages are readable')
bullet('mock-frontend/index.html — home page with links to both forms')

h2('tests/')
p(
    'All test files live here, split into three sub-folders.'
)
bullet(
    'tests/api/ — API test specs. '
    'users.api.spec.ts tests the user endpoints; '
    'transactions.api.spec.ts tests the transaction endpoints. '
    'These tests do not open a browser — they send HTTP requests directly.'
)
bullet(
    'tests/ui/ — Browser test specs. '
    'registration.ui.spec.ts drives the registration form in Chrome; '
    'transaction.ui.spec.ts drives the transaction form.'
)
bullet(
    'tests/utils/ — Shared utilities used by all the tests. '
    'factories.ts builds test data, '
    'apiClient.ts wraps the HTTP calls, '
    'assertions.ts provides readable pass/fail helpers, '
    'logger.ts saves API responses to disk for debugging.'
)

h2('Root-level files')
bullet('playwright.config.ts — tells Playwright where the tests are, which browser to use, what reports to generate, and how to start the mock server')
bullet('package.json — lists the project dependencies and defines the npm run commands')
bullet('tsconfig.json — TypeScript settings: tells the compiler which language features to allow and where to put the output')
bullet('README.md — setup instructions, design notes, and a description of what was skipped')
bullet('.gitignore — tells Git not to track the node_modules folder, build output, or test reports')

# ── What was skipped ──────────────────────────────────────────────────────────
h1('What Was Not Included and Why')
p(
    'The Notification Service was left out because the brief did not specify an endpoint '
    'for it. With more time I would add:'
)
bullet('Notification service tests once the endpoint contract is known')
bullet('Data-driven tests for the validation cases — testing many combinations of bad input from a table rather than one test per case')
bullet('A GitHub Actions workflow file so the suite runs automatically on every pull request')
bullet('Load/performance tests for the transaction endpoint using a tool like k6')

doc.save('QE_Framework_Submission.docx')
print('Saved QE_Framework_Submission.docx')
